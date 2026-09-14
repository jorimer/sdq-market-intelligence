"""OC-SENI connector — Informe Mensual de Transacciones Económicas (IMTE), Eje Energía.

El **Organismo Coordinador del Sistema Eléctrico Nacional Interconectado** (OC-SENI) publica
cada mes, entre el día 20 y el 22, el IMTE del mes anterior: un ZIP de ~80 MB con los balances
del mercado eléctrico mayorista. De ese ZIP se lee UN libro —«08 09 Resumen Transacciones e
Indicadores <año>.xlsx»— y de él UNA hoja, «Iny», con la energía del sistema por mes.

**La API es pública y sin login** (``apps.oc.org.do/website-api/v1``), verificada end-to-end el
2026-09-07 y otra vez el 2026-09-14:

* Listado: ``GET publicaciones/<Área>?loggedIn=false`` → ``[{id, ruta, nombre_archivo, tamano,
  fecha_publicacion}]``. La ``ruta`` nombra el MES OPERADO (``…/2026/07. Jul/Versión 0``); la
  ``fecha_publicacion`` es cuándo lo publicó el OC, y es la fecha que el delta cita.
* Descarga: ``POST publicaciones/download-url`` con ``{"files_id": [id]}`` → ``[url]``. **La URL
  es de un solo uso**: un segundo GET devuelve 404, y un pedido con ``Range`` consume el token
  bajando el archivo entero. Se lee en un único GET, en streaming.

⚠️ **Un área inexistente devuelve 200 con ``[]``, no un error.** Un typo en el nombre del área se
leería como «el OC no publicó nada» — y el sensor de fuentes congeladas lo daría por fuente
parada. Por eso las áreas pasan por una LISTA BLANCA y fuera de ella se lanza. (Con espacios en
vez de guiones, además, devuelve 500.)

**El orden de los bloques de año se VERIFICA, no se supone.** La hoja trae dos bloques con los
doce meses: el del año del libro y, debajo, el del año anterior. Que el de arriba sea el año en
curso se comprueba contra una cifra que el propio OC publica: la «Variación Inyecciones (respecto
mes anterior)» de enero tiene que ser ``ENE(año) / DIC(año-1) − 1``. Si no coincide, el parseo
falla: los dos bloques invertidos darían un año de movimiento con el signo cambiado.

**Qué NO se lee.** El costo marginal (hoja «Indicadores Costos Marginales») es un PRECIO: con la
unidad en RD$ la inferencia de naturaleza lo tomaría por flujo y la ventana de doce meses lo
sumaría. Queda fuera hasta declararlo.
"""
from __future__ import annotations

import io
import json
import logging
import re
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date
from typing import Any, Dict, List, Optional, Sequence, Tuple

from shared.data.base_client import _FIXTURES_DIR, Mode, Record, SourceClient
from shared.data.lineage import Lineage

logger = logging.getLogger("sdq.data.oc_seni")

_API = "https://apps.oc.org.do/website-api/v1/publicaciones/"
_HEADERS = {"User-Agent": "Mozilla/5.0 (SDQ-MIP)"}

#: El área del sitio del OC donde vive el IMTE. Va con guiones y con tildes, tal cual la API.
AREA_TRANSACCIONES = "Transacciones-Económicas-y-Cálculos-Comerciales"

#: Las únicas áreas que este conector pide. Una fuera de acá LANZA: la API contesta 200 con una
#: lista vacía a cualquier nombre que no existe.
AREAS_CONOCIDAS = frozenset({AREA_TRANSACCIONES})

SERIE_INYECCIONES = "oc_seni.imte.inyecciones_gwh"
SERIE_RETIROS = "oc_seni.imte.retiros_totales_gwh"
SERIE_RETIROS_DISTRIBUIDORAS = "oc_seni.imte.retiros_distribuidoras_gwh"
SERIE_PERDIDAS = "oc_seni.imte.perdidas_transmision_pct"

#: Etiqueta de la fila en la hoja «Iny» (normalizada) → (serie, unidad). «% Pérdidas Totales» es
#: la energía inyectada que no se retira —(inyecciones − retiros) / inyecciones, verificado contra
#: la hoja—, o sea pérdidas de TRANSMISIÓN del sistema, no las de distribución de las EDE.
FILAS_IMTE: Dict[str, Tuple[str, str]] = {
    "inyecciones (gwh)": (SERIE_INYECCIONES, "GWh"),
    "retiros totales (gwh)": (SERIE_RETIROS, "GWh"),
    "retiros distribuidoras (gwh)": (SERIE_RETIROS_DISTRIBUIDORAS, "GWh"),
    "% perdidas totales": (SERIE_PERDIDAS, "%"),
}
_FILA_VARIACION = "variacion inyecciones"

_MESES = ("ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic")
_RE_EDICION = re.compile(r"^OC-GC-07-IMTE-(\d{8})-V(\d+)\.zip$", re.IGNORECASE)
_RE_RUTA = re.compile(r"/(\d{4})/(\d{2})\.\s*[^/]*/Versi[oó]n\s*(\d+)", re.IGNORECASE)
_RE_MIEMBRO = re.compile(r"Resumen Transacciones e Indicadores (\d{4})\.xlsx$", re.IGNORECASE)
_TOLERANCIA_VARIACION = 1e-6


class AreaDesconocida(ValueError):
    """Un área del OC que no está en la lista blanca: la API la contestaría con ``[]``."""


class EstructuraInesperada(ValueError):
    """La hoja no tiene la forma verificada. Se falla: adivinar publicaría cifras corridas."""


@dataclass(frozen=True)
class EdicionIMTE:
    """Una edición del IMTE en el listado del OC."""

    file_id: str
    nombre_archivo: str
    periodo: str            # el mes OPERADO, "YYYY-MM"
    version: int
    publicada: Optional[date]


def _norm(v: Any) -> str:
    t = "".join(c for c in unicodedata.normalize("NFKD", str(v).lower())
                if not unicodedata.combining(c))
    return " ".join(t.split())


def validar_area(area: str) -> str:
    if area not in AREAS_CONOCIDAS:
        raise AreaDesconocida(
            f"OC-SENI: el área {area!r} no está en la lista blanca {sorted(AREAS_CONOCIDAS)}. "
            f"La API devuelve 200 con [] para un área inexistente: pedirla se leería como que "
            f"el OC no publicó nada.")
    return area


def url_del_listado(area: str) -> str:
    from urllib.parse import quote

    return _API + quote(validar_area(area), safe="-") + "?loggedIn=false"


def ediciones_imte(listado: Sequence[Dict[str, Any]]) -> List[EdicionIMTE]:
    """Las ediciones del IMTE del listado, UNA por mes operado (la de versión más alta).

    El mes sale de la ``ruta`` y no del nombre del archivo: el nombre lleva la fecha de
    PUBLICACIÓN (``20260820`` es la edición de julio).
    """
    por_mes: Dict[str, EdicionIMTE] = {}
    for item in listado:
        nombre = str(item.get("nombre_archivo") or "")
        m_nombre = _RE_EDICION.match(nombre)
        m_ruta = _RE_RUTA.search(str(item.get("ruta") or ""))
        if not m_nombre or not m_ruta:
            continue
        publicada = None
        try:
            publicada = date.fromisoformat(str(item.get("fecha_publicacion") or "")[:10])
        except ValueError:
            pass
        ed = EdicionIMTE(file_id=str(item["id"]), nombre_archivo=nombre,
                         periodo=f"{m_ruta.group(1)}-{m_ruta.group(2)}",
                         version=int(m_nombre.group(2)), publicada=publicada)
        previa = por_mes.get(ed.periodo)
        if previa is None or ed.version > previa.version:
            por_mes[ed.periodo] = ed
    return [por_mes[p] for p in sorted(por_mes)]


def _numero(v: Any) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None   # un texto en una celda de cifra no es una cifra


def parse_hoja_iny(filas: Sequence[Sequence[Any]], anio: int,
                   hasta: str) -> Dict[str, Dict[str, Optional[float]]]:
    """``{serie: {"YYYY-MM": valor}}`` de la hoja «Iny», del año anterior hasta *hasta*.

    Los meses posteriores a *hasta* no se leen: la hoja los trae vacíos o en cero, y un cero en
    un mes que no ocurrió se publicaría como una medición.
    """
    bloques: List[Tuple[int, Dict[str, int]]] = []
    for i, fila in enumerate(filas):
        idx = {_norm(v): j for j, v in enumerate(fila) if v is not None and _norm(v) in _MESES}
        if len(idx) == len(_MESES):
            bloques.append((i, idx))
    if len(bloques) != 2:
        raise EstructuraInesperada(
            f"OC-SENI: la hoja «Iny» tiene {len(bloques)} bloques de meses; se esperaban 2 "
            f"(el año del libro y el anterior).")

    def filas_del_bloque(nb: int) -> Dict[str, Dict[str, Any]]:
        inicio, idx = bloques[nb]
        fin = bloques[nb + 1][0] if nb + 1 < len(bloques) else len(filas)
        out: Dict[str, Dict[str, Any]] = {}
        for fila in filas[inicio + 1:fin]:
            etiqueta = next((v for v in fila[:idx["ene"]] if v not in (None, "")), None)
            if etiqueta is None:
                continue
            out[_norm(etiqueta)] = {m: (fila[j] if j < len(fila) else None)
                                    for m, j in idx.items()}
        return out

    actual, anterior = filas_del_bloque(0), filas_del_bloque(1)
    faltan = [e for e in (*FILAS_IMTE, "inyecciones (gwh)") if e not in actual or e not in anterior]
    variacion = next((v for k, v in actual.items() if k.startswith(_FILA_VARIACION)), None)
    if faltan or variacion is None:
        raise EstructuraInesperada(
            f"OC-SENI: faltan filas en la hoja «Iny»: {sorted(set(faltan)) or ['variación']}.")

    # EL ORDEN DE LOS BLOQUES SE VERIFICA contra una cifra del propio OC.
    ene, dic_previo = (_numero(actual["inyecciones (gwh)"]["ene"]),
                       _numero(anterior["inyecciones (gwh)"]["dic"]))
    publicada = _numero(variacion["ene"])
    if ene is None or not dic_previo or publicada is None:
        raise EstructuraInesperada("OC-SENI: sin enero, diciembre previo o variación publicada "
                                   "no se puede verificar el orden de los años.")
    if abs((ene / dic_previo - 1) - publicada) > _TOLERANCIA_VARIACION:
        raise EstructuraInesperada(
            f"OC-SENI: ENE/DIC−1 = {ene / dic_previo - 1:.6f} y el OC publica {publicada:.6f}. "
            f"Los bloques de año no están en el orden verificado.")

    series: Dict[str, Dict[str, Optional[float]]] = {s: {} for s, _ in FILAS_IMTE.values()}
    for bloque, anio_bloque in ((anterior, anio - 1), (actual, anio)):
        for etiqueta, (serie, unidad) in FILAS_IMTE.items():
            for n, mes in enumerate(_MESES, start=1):
                periodo = f"{anio_bloque}-{n:02d}"
                if periodo > hasta:
                    continue
                valor = _numero(bloque[etiqueta][mes])
                if valor is not None and unidad == "%":
                    valor = valor * 100.0   # la hoja lo trae como fracción
                series[serie][periodo] = valor
    return series


def parse_libro_resumen(contenido: bytes, anio: int, hasta: str) -> Dict[str, Dict[str, Optional[float]]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    try:
        if "Iny" not in wb.sheetnames:
            raise EstructuraInesperada(f"OC-SENI: el libro no tiene hoja «Iny» ({wb.sheetnames}).")
        filas = list(wb["Iny"].iter_rows(min_row=1, max_row=60, values_only=True))
    finally:
        wb.close()
    return parse_hoja_iny(filas, anio, hasta)


def libro_resumen_del_zip(archivo: Any) -> Tuple[bytes, int]:
    """``(bytes del libro resumen, año del libro)`` dentro del ZIP de una edición."""
    with zipfile.ZipFile(archivo) as zf:
        for nombre in zf.namelist():
            m = _RE_MIEMBRO.search(nombre)
            if m:
                return zf.read(nombre), int(m.group(1))
    raise EstructuraInesperada("OC-SENI: el ZIP del IMTE no trae el libro «Resumen "
                               "Transacciones e Indicadores <año>.xlsx».")


class OCSENIClient(SourceClient):
    source = "OC-SENI (Organismo Coordinador del Sistema Eléctrico Nacional Interconectado)"
    license = ("OC-SENI — informe mensual de transacciones económicas (IMTE) publicado sin login "
               "en apps.oc.org.do. Los Términos de uso del sitio (plantilla del CMS) no "
               "restringen la reutilización del dato; el IMTE es una publicación que el OC hace "
               "en cumplimiento de la Ley General de Electricidad 125-01. Reutilizable con "
               "atribución.")
    license_ok = True
    fixture_file = "oc_seni_imte_resumen_2026.xlsx"
    #: La edición que describe el fixture, para que el modo fixture diga de qué mes es.
    fixture_edicion = EdicionIMTE(file_id="fixture", nombre_archivo="OC-GC-07-IMTE-20260820-V0.zip",
                                  periodo="2026-07", version=0, publicada=date(2026, 8, 20))

    def __init__(self, mode: Mode = "fixture") -> None:
        super().__init__(mode)

    def listar(self, area: str = AREA_TRANSACCIONES) -> List[Dict[str, Any]]:
        import httpx

        url = url_del_listado(area)
        with httpx.Client(timeout=90, follow_redirects=True, headers=_HEADERS) as http:
            r = http.get(url)
            r.raise_for_status()
            datos = r.json()
        if not isinstance(datos, list):
            raise EstructuraInesperada(f"OC-SENI: el listado no es una lista ({type(datos).__name__}).")
        return datos

    def _descargar(self, edicion: EdicionIMTE) -> Any:
        """El ZIP en un archivo temporal. UN solo GET: la URL firmada no admite un segundo."""
        import httpx

        with httpx.Client(timeout=120, follow_redirects=True, headers=_HEADERS) as http:
            r = http.post(_API + "download-url", json={"files_id": [edicion.file_id]})
            r.raise_for_status()
            urls = r.json()
            if not isinstance(urls, list) or not urls:
                raise EstructuraInesperada(f"OC-SENI: download-url no devolvió una URL ({urls!r}).")
            tmp = tempfile.TemporaryFile()
            with http.stream("GET", str(urls[0]), timeout=900) as s:
                s.raise_for_status()
                for trozo in s.iter_bytes(1 << 20):
                    tmp.write(trozo)
        tmp.seek(0)
        return tmp

    def fetch(self, series: Optional[str] = None, period: Optional[str] = None) -> List[Record]:
        """El contrato de `SourceClient`: los registros, filtrados. La ingesta usa
        :meth:`leer_ultima_edicion`, que además dice de qué edición salen."""
        registros, _ = self.leer_ultima_edicion()
        return [r for r in registros
                if (series is None or r.series == series) and (period is None or r.period == period)]

    def leer_ultima_edicion(self) -> Tuple[List[Record], EdicionIMTE]:
        """Las series del IMTE de la ÚLTIMA edición publicada, con la edición de la que salen."""
        self.check_license()
        if self.mode == "fixture":
            edicion = self.fixture_edicion
            contenido, anio = (_FIXTURES_DIR / self.fixture_file).read_bytes(), 2026
        else:
            ediciones = ediciones_imte(self.listar())
            if not ediciones:
                raise EstructuraInesperada("OC-SENI: el listado no trae ediciones del IMTE.")
            edicion = ediciones[-1]
            with self._descargar(edicion) as archivo:
                contenido, anio = libro_resumen_del_zip(archivo)
        series = parse_libro_resumen(contenido, anio, edicion.periodo)
        lineage = Lineage(source=self.source, license=self.license, fetched_at=date.today())
        unidades = dict(FILAS_IMTE.values())
        registros = [Record(series=s, period=p, value=v, lineage=lineage, unit=unidades[s])
                     for s, puntos in series.items() for p, v in sorted(puntos.items())]
        logger.info("[OC-SENI] IMTE %s (%s): %d puntos", edicion.periodo, self.mode, len(registros))
        return registros, edicion


__all__ = ["AREA_TRANSACCIONES", "AREAS_CONOCIDAS", "AreaDesconocida", "EdicionIMTE",
           "EstructuraInesperada", "FILAS_IMTE", "OCSENIClient", "SERIE_INYECCIONES",
           "SERIE_PERDIDAS", "SERIE_RETIROS", "SERIE_RETIROS_DISTRIBUIDORAS", "ediciones_imte",
           "json", "libro_resumen_del_zip", "parse_hoja_iny", "parse_libro_resumen",
           "url_del_listado", "validar_area"]
